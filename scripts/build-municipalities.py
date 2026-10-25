#!/usr/bin/env python3
"""Offline projection of the ISTAT XLSX snapshot. Requires openpyxl for maintainers only.
Usage: python3 scripts/build-municipalities.py /path/Elenco-comuni-italiani.xlsx
Download source and review its reference date before refreshing the checked-in JSON.
"""
from pathlib import Path
from openpyxl import load_workbook
import argparse, hashlib, json
parser=argparse.ArgumentParser();parser.add_argument('xlsx',type=Path);args=parser.parse_args()
raw=args.xlsx;sheet=load_workbook(raw,read_only=True,data_only=True).active
rows=[[str(r[4]).zfill(6),r[6],r[14] or '',r[10],r[5]] for r in sheet.iter_rows(min_row=2,values_only=True) if r[4]]
assert len(rows)==7894 and len({r[0] for r in rows})==7894,'Review updated snapshot/count before refreshing'
obj={'schemaVersion':1,'source':'ISTAT','sourceURL':'https://www.istat.it/storage/codici-unita-amministrative/Elenco-comuni-italiani.xlsx','referenceDate':'2026-02-21','retrievedDate':'2026-10-25','sourceSHA256':hashlib.sha256(raw.read_bytes()).hexdigest(),'fields':['istat','name','province','region','officialName'],'municipalities':sorted(rows,key=lambda r:r[1].casefold()),'license':'CC BY 4.0','licenseURL':'https://creativecommons.org/licenses/by/4.0/','transformation':'Selected administrative columns from ISTAT XLSX, strings preserved and rows sorted by Italian name. No market data or geographic inference.'}
output=Path(__file__).resolve().parents[1]/'data/geography/municipalities-it.json';output.parent.mkdir(parents=True,exist_ok=True);output.write_text(json.dumps(obj,ensure_ascii=False,separators=(',',':'))+'\n');print(f'{len(rows)} municipalities written')
